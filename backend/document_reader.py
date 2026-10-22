import io
from pathlib import Path

import fitz
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from PIL import Image, ImageOps


# ============================================================
# SUPPORTED FILE TYPES
# ============================================================

IMAGE_EXTS = {
    'jpg',
    'jpeg',
    'png'
}

SUPPORTED = {
    'pdf',
    'doc',
    'docx',
    'ppt',
    'pptx',
    'xlsx',
    'txt',
    *IMAGE_EXTS
}


# ============================================================
# IMAGE NORMALIZATION
# ============================================================

def normalize_image(data: bytes) -> bytes:
    """
    Convert uploaded images to RGB JPEG and resize them
    to a reasonable size before sending them to Bedrock.
    """

    with Image.open(io.BytesIO(data)) as im:

        # Correct mobile-camera EXIF rotation
        im = ImageOps.exif_transpose(im)

        # Nova receives RGB JPEG
        im = im.convert('RGB')

        # Prevent unnecessarily huge images
        im.thumbnail((1800, 1800))

        out = io.BytesIO()

        im.save(
            out,
            'JPEG',
            quality=83,
            optimize=True
        )

        return out.getvalue()


# ============================================================
# MAIN DOCUMENT READER
# ============================================================

def read_document(filename: str, data: bytes):

    ext = (
        Path(filename)
        .suffix
        .lower()
        .lstrip('.')
    )

    # --------------------------------------------------------
    # CHECK FILE TYPE
    # --------------------------------------------------------

    if ext not in SUPPORTED:
        raise ValueError(
            'Unsupported file extension'
        )


    # ========================================================
    # JPG / JPEG / PNG
    # ========================================================

    if ext in IMAGE_EXTS:

        return {
            'text': '',
            'images': [
                normalize_image(data)
            ],
            'warnings': []
        }


    # ========================================================
    # TXT
    # ========================================================

    if ext == 'txt':

        return {
            'text': data.decode(
                'utf-8-sig',
                errors='replace'
            ),
            'images': [],
            'warnings': []
        }


    # ========================================================
    # PDF
    # ========================================================

    if ext == 'pdf':

        pdf = PdfReader(
            io.BytesIO(data)
        )

        if pdf.is_encrypted:
            raise ValueError(
                'Encrypted PDF is not supported'
            )

        # ----------------------------------------------------
        # First try normal PDF text extraction
        # ----------------------------------------------------

        text = '\n\n'.join(
            (page.extract_text() or '')
            for page in pdf.pages
        )

        # ----------------------------------------------------
        # NORMAL TEXT-BASED PDF
        # ----------------------------------------------------

        if text.strip():

            return {
                'text': text,
                'images': [],
                'warnings': []
            }

        # ----------------------------------------------------
        # SCANNED PDF
        #
        # pypdf cannot read text because the PDF consists
        # of page images.
        #
        # Convert the pages to images and allow Amazon
        # Nova 2 Lite to visually analyse them.
        # ----------------------------------------------------

        pdf_doc = fitz.open(
            stream=data,
            filetype='pdf'
        )

        images = []

        # Prototype limit:
        # send a maximum of 10 pages to Bedrock
        page_count = min(
            len(pdf_doc),
            10
        )

        for page_number in range(page_count):

            page = pdf_doc[
                page_number
            ]

            # Render PDF page at improved resolution
            pix = page.get_pixmap(
                matrix=fitz.Matrix(
                    1.5,
                    1.5
                ),
                alpha=False
            )

            # Convert rendered page to PNG bytes
            png_bytes = pix.tobytes(
                'png'
            )

            # Normalize it to JPEG for Bedrock
            image_bytes = normalize_image(
                png_bytes
            )

            images.append(
                image_bytes
            )

        pdf_doc.close()

        if not images:
            raise ValueError(
                'No readable content found in PDF.'
            )

        warnings = [
            'Scanned PDF detected. '
            'PDF pages were converted to images '
            'for Amazon Nova 2 Lite analysis.'
        ]

        if len(pdf.pages) > 10:

            warnings.append(
                'This prototype processes only the '
                'first 10 pages of a scanned PDF.'
            )

        return {
            'text': '',
            'images': images,
            'warnings': warnings
        }


    # ========================================================
    # DOCX
    # ========================================================

    if ext == 'docx':

        doc = Document(
            io.BytesIO(data)
        )

        parts = []

        # Extract paragraphs
        for paragraph in doc.paragraphs:

            if paragraph.text.strip():

                parts.append(
                    paragraph.text
                )

        # Extract tables
        for table in doc.tables:

            for row in table.rows:

                row_text = ' | '.join(
                    cell.text
                    for cell in row.cells
                )

                if row_text.strip():

                    parts.append(
                        row_text
                    )

        return {
            'text': '\n'.join(parts),
            'images': [],
            'warnings': []
        }


    # ========================================================
    # PPTX
    # ========================================================

    if ext == 'pptx':

        pres = Presentation(
            io.BytesIO(data)
        )

        parts = []

        for slide_number, slide in enumerate(
            pres.slides,
            start=1
        ):

            parts.append(
                f'Slide {slide_number}'
            )

            for shape in slide.shapes:

                # Extract normal text
                if shape.has_text_frame:

                    text = shape.text.strip()

                    if text:

                        parts.append(
                            text
                        )

                # Extract PowerPoint tables
                if shape.has_table:

                    for row in shape.table.rows:

                        row_text = ' | '.join(
                            cell.text
                            for cell in row.cells
                        )

                        if row_text.strip():

                            parts.append(
                                row_text
                            )

        return {
            'text': '\n'.join(parts),
            'images': [],
            'warnings': []
        }


    # ========================================================
    # XLSX
    # ========================================================

    if ext == 'xlsx':

        book = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True
        )

        parts = []

        for sheet in book:

            parts.append(
                f'Sheet: {sheet.title}'
            )

            for row in sheet.iter_rows(
                values_only=True
            ):

                if any(
                    value is not None
                    for value in row
                ):

                    row_text = ' | '.join(
                        ''
                        if value is None
                        else str(value)
                        for value in row
                    )

                    parts.append(
                        row_text
                    )

        book.close()

        return {
            'text': '\n'.join(parts),
            'images': [],
            'warnings': []
        }


    # ========================================================
    # LEGACY DOC / PPT
    # ========================================================

    if ext in {
        'doc',
        'ppt'
    }:

        raise ValueError(
            'Legacy DOC and PPT files cannot be read '
            'directly by this prototype. '
            'Convert them to DOCX or PPTX first.'
        )


    # ========================================================
    # FALLBACK
    # ========================================================

    raise ValueError(
        'Unable to process this document.'
    )